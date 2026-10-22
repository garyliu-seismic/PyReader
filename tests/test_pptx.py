import io
import pytest

pptx = pytest.importorskip("pptx")
from pptx import Presentation
from pptx.util import Inches

from pyreader.pptx import load_pptx

NOTE = "SECRET-SPEAKER-NOTE"


def _png():
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (8, 6), (255, 0, 0)).save(buf, "PNG")
    return buf.getvalue()


def build(path):
    prs = Presentation()
    # 1: 标题 + 项目符号（含二级）+ 备注
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Intro"
    tf = s.placeholders[1].text_frame
    tf.text = "first point"
    p = tf.add_paragraph()
    p.text = "nested point"
    p.level = 1
    s.notes_slide.notes_text_frame.text = NOTE
    # 2: 表格
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Data"
    tbl = s.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    for r, row in enumerate([["a", "b"], ["c", "d"]]):
        for c, v in enumerate(row):
            tbl.cell(r, c).text = v
    # 3: 空幻灯片（无标题）
    prs.slides.add_slide(prs.slide_layouts[6])
    # 4: 分组形状 + 阅读顺序（下方的先加入）+ 图片
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Shapes"
    s.shapes.add_textbox(Inches(1), Inches(5), Inches(2), Inches(1)).text_frame.text = "bottom"
    g = s.shapes.add_group_shape()
    g.shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(1)).text_frame.text = "in group"
    s.shapes.add_textbox(Inches(1), Inches(2), Inches(2), Inches(1)).text_frame.text = "top"
    s.shapes.add_textbox(Inches(1), Inches(6), Inches(2), Inches(1)).text_frame.text = ""
    s.shapes.add_picture(io.BytesIO(_png()), Inches(1), Inches(1.5))
    prs.save(str(path))
    return str(path)


@pytest.fixture
def deck(tmp_path):
    return load_pptx(build(tmp_path / "t.pptx"))


def test_chapters(deck):
    text, chapters, title, extra = deck
    assert [t for _, t in chapters] == ["Intro", "Data", "Slide 3", "Shapes"]
    for off, t in chapters:
        assert text[off:off + len(t)] == t


def test_block_starts_match_text(deck):
    text, _, _, extra = deck
    for b in extra["blocks"]:
        if b["type"] in ("heading", "para", "list"):
            s = "".join(p[0] for p in b["pieces"])
            assert text[b["start"]:b["start"] + len(s)] == s
            for t, _bits, a, e in b["pieces"]:
                assert text[a:e] == t
        elif b["type"] == "table":
            for row in b["cells"]:
                for c, off in row:
                    assert text[off:off + len(c)] == c


def test_bullets_and_levels(deck):
    lists = [b for b in deck[3]["blocks"] if b["type"] == "list"]
    assert [(b["pieces"][0][0], b["depth"]) for b in lists] == [
        ("first point", 0), ("nested point", 1)]


def test_table(deck):
    tbls = [b for b in deck[3]["blocks"] if b["type"] == "table"]
    assert len(tbls) == 1
    assert [[c for c, _ in r] for r in tbls[0]["cells"]] == [["a", "b"], ["c", "d"]]


def test_empty_slide_has_fallback_heading_only(deck):
    blocks = deck[3]["blocks"]
    i = next(i for i, b in enumerate(blocks)
             if b["type"] == "heading" and b["pieces"][0][0] == "Slide 3")
    assert blocks[i + 1]["type"] == "heading"   # 下一块直接是下一张的标题


def test_reading_order_group_and_empty_skipped(deck):
    text = deck[0]
    # 标题 → 顶部图片区域 → top → 组内 → bottom（空文本框被跳过）
    assert text.index("Shapes") < text.index("top") < text.index("in group") < text.index("bottom")
    shapes_blocks = deck[3]["blocks"][-5:]
    assert [b["type"] for b in shapes_blocks] == ["heading", "image", "para", "para", "para"]


def test_image(deck):
    imgs = [b for b in deck[3]["blocks"] if b["type"] == "image"]
    assert len(imgs) == 1
    b = imgs[0]
    assert deck[3]["images"][b["img_id"]][:4] == b"\x89PNG"
    assert (b["img_w"], b["img_h"]) == (8, 6)


def test_notes_not_extracted(deck):
    assert NOTE not in deck[0]
    assert NOTE not in str(deck[3]["blocks"])


def test_loader_dispatch(tmp_path):
    pytest.importorskip("PySide6")
    from pyreader.loader import BookLoader
    assert ".pptx" in BookLoader._FULL_PARSE_EXTS
    got = {}
    bl = BookLoader(build(tmp_path / "d.pptx"))
    bl.loaded.connect(lambda *a: got.update(r=a))
    bl.failed.connect(lambda m: got.update(err=m))
    bl.run()
    assert "err" not in got, got.get("err")
    text, chapters, _title, extra = got["r"]
    assert [t for _, t in chapters][0] == "Intro"
    assert extra["blocks"] and NOTE not in text


# ---------------- 背景填充 ----------------
from lxml import etree
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT


def _png_color(c):
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (8, 6), c).save(buf, "PNG")
    return buf.getvalue()


def _set_pic_bg(owner, blob_bytes):
    """给 slide / layout / master 设置图片背景（python-pptx 无 API，手写 XML）。"""
    from pptx.parts.image import ImagePart
    part = owner.part
    img_part = ImagePart.new(part.package, __import__("pptx.parts.image", fromlist=["Image"]).Image.from_blob(blob_bytes))
    rid = part.relate_to(img_part, RT.IMAGE)
    csld = owner._element.find(qn("p:cSld"))
    old = csld.find(qn("p:bg"))
    if old is not None:
        csld.remove(old)
    bg = etree.fromstring(
        '<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<p:bgPr><a:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></a:blipFill>'
        '<a:effectLst/></p:bgPr></p:bg>')
    csld.insert(0, bg)


def _set_solid_bg(owner):
    s = owner.background.fill
    s.solid()
    from pptx.dml.color import RGBColor
    s.fore_color.rgb = RGBColor(0, 0, 255)


def _bg_blocks(extra):
    return [b for b in extra["blocks"] if b["type"] == "image"]


def test_bg_picture_on_slide_follows_heading(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "T"
    _set_pic_bg(s, _png_color((1, 2, 3)))
    prs.save(str(tmp_path / "a.pptx"))
    _, _, _, extra = load_pptx(str(tmp_path / "a.pptx"))
    assert [b["type"] for b in extra["blocks"]] == ["heading", "image"]
    assert extra["blocks"][1]["start"] == extra["blocks"][0]["start"] + 2
    assert (extra["blocks"][1]["img_w"], extra["blocks"][1]["img_h"]) == (8, 6)


def test_bg_inherited_from_layout_and_master_stored_once(tmp_path):
    prs = Presentation()
    lay = prs.slide_layouts[5]
    _set_pic_bg(lay, _png_color((9, 9, 9)))
    for i in range(3):
        prs.slides.add_slide(lay).shapes.title.text = f"L{i}"
    m = prs.slide_layouts[6]            # 该版式无背景 → 继承母版
    _set_pic_bg(prs.slide_master, _png_color((7, 7, 7)))
    prs.slides.add_slide(m)
    prs.save(str(tmp_path / "b.pptx"))
    _, _, _, extra = load_pptx(str(tmp_path / "b.pptx"))
    imgs = _bg_blocks(extra)
    assert len(imgs) == 4                       # 每张幻灯片一个块
    assert len(extra["images"]) == 2            # 版式图 + 母版图各存一份
    assert len({b["img_id"] for b in imgs[:3]}) == 1


def test_bg_solid_no_image_and_blocks_inheritance(tmp_path):
    prs = Presentation()
    lay = prs.slide_layouts[5]
    _set_pic_bg(lay, _png_color((9, 9, 9)))
    s = prs.slides.add_slide(lay)
    s.shapes.title.text = "T"
    _set_solid_bg(s)                            # 自身纯色覆盖继承的图片
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    _set_solid_bg(s2)
    prs.save(str(tmp_path / "c.pptx"))
    _, _, _, extra = load_pptx(str(tmp_path / "c.pptx"))
    assert _bg_blocks(extra) == [] and extra["images"] == {}


# ---------------- 仅图片背景的幻灯片 ----------------
def _image_only_deck(tmp_path):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Before"
    s2 = prs.slides.add_slide(prs.slide_layouts[6])      # 无标题无形状，仅背景图
    _set_pic_bg(s2, _png_color((5, 5, 5)))
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "After"
    _set_pic_bg(s3, _png_color((6, 6, 6)))               # 有标题 + 背景图，其后无文字
    prs.save(str(tmp_path / "io.pptx"))
    return load_pptx(str(tmp_path / "io.pptx"))


def test_image_only_slide_no_visible_heading(tmp_path):
    text, chapters, _, extra = _image_only_deck(tmp_path)
    assert "Slide 2" not in text
    assert [t for _, t in chapters] == ["Before", "Slide 2", "After"]
    blocks = extra["blocks"]
    assert [b["type"] for b in blocks] == ["heading", "image", "heading", "image"]
    # 图片块属于自己的幻灯片：恰好落在该章节范围内，而不是下一章的起点
    assert chapters[1][0] == blocks[1]["start"]
    assert chapters[1][0] <= blocks[1]["start"] < chapters[2][0]
    assert chapters[2][0] <= blocks[3]["start"] < len(text)
    assert text[blocks[1]["start"]] == "\n"
    # 图片占 1 个字符：块 start 单调递增且互不相同
    starts = [b["start"] for b in blocks]
    assert starts == sorted(set(starts))
    assert chapters[0][0] == 0 and text.startswith("Before\n")
