"""PPTX 解析 → (全文, 章节, 书名, 富文本块)，结构与 docx.load_docx 一致。

每张幻灯片 = 一个 heading 块（幻灯片标题，无标题则 "Slide N"）+ 正文段落/列表/表格/图片。
阅读顺序：标题占位符优先，其余形状按 (top, left) 排序，组合形状递归展开。
演讲者备注刻意不提取（朗读时不应读出备注）。"""
import io

from .markdown import MD_BOLD, MD_ITALIC

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _sort_key(shape):
    return (shape.top if shape.top is not None else 0,
            shape.left if shape.left is not None else 0)


def _ph_type(shape):
    """占位符类型（非占位符返回 None）。"""
    try:
        if shape.is_placeholder:
            return shape.placeholder_format.type
    except Exception:
        pass
    return None


def _is_title(shape):
    from pptx.enum.shapes import PP_PLACEHOLDER as PH
    return _ph_type(shape) in (PH.TITLE, PH.CENTER_TITLE, PH.VERTICAL_TITLE)


def _para_pieces(p, start):
    """段落 → pieces，元素 (text, bits, start, end)，含粗斜体位。"""
    pieces, local = [], start
    for child in p._p:
        tag = child.tag
        if tag == _A + "br":
            t, bits = " ", 0
        elif tag in (_A + "r", _A + "fld"):
            te = child.find(_A + "t")
            t = (te.text or "") if te is not None else ""
            rpr = child.find(_A + "rPr")
            bits = 0
            if rpr is not None:
                if rpr.get("b") in ("1", "true"):
                    bits |= MD_BOLD
                if rpr.get("i") in ("1", "true"):
                    bits |= MD_ITALIC
        else:
            continue
        if not t:
            continue
        pieces.append((t, bits, local, local + len(t)))
        local += len(t)
    return pieces


def _bullet_kind(p, in_body):
    """返回 None（普通段落）/ "bullet" / "num"。"""
    ppr = p._p.find(_A + "pPr")
    if ppr is not None:
        if ppr.find(_A + "buNone") is not None:
            return None
        if ppr.find(_A + "buAutoNum") is not None:
            return "num"
        if ppr.find(_A + "buChar") is not None:
            return "bullet"
    if in_body:
        return "bullet"
    return "bullet" if p.level > 0 else None


def _iter_shapes(shapes):
    """按阅读顺序展开形状（组合形状递归）。"""
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    for sh in sorted(shapes, key=_sort_key):
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(sh.shapes)
        else:
            yield sh


def load_pptx(path, progress=None):
    """解析 PPTX → (全文, 章节, 书名, extra)；extra = {"blocks": [...], "images": {...}}。"""
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        from pptx.enum.shapes import PP_PLACEHOLDER as PH
    except Exception as e:
        raise RuntimeError("未安装 python-pptx，请执行：pip install python-pptx") from e

    if progress:
        progress(5, "解析 pptx…")

    prs = Presentation(path)
    title = (prs.core_properties.title or "").strip()

    blocks, images, chapters, text_parts = [], {}, [], []
    offset = 0
    img_counter = 0

    def emit_text(ttype, pieces, **kw):
        nonlocal offset
        start = offset
        txt = "".join(p[0] for p in pieces)
        text_parts.append(txt)
        text_parts.append("\n")
        blk = {"type": ttype, "start": start, "pieces": pieces}
        blk.update(kw)
        blocks.append(blk)
        offset += len(txt) + 1
        return start, txt

    def emit_table(table):
        nonlocal offset
        start = offset
        rows = [[cell.text.replace("\n", " ").replace("\x0b", " ") for cell in row.cells]
                for row in table.rows]
        if not any(c.strip() for r in rows for c in r):
            return
        txt = "\n".join("\t".join(r) for r in rows)
        text_parts.append(txt)
        text_parts.append("\n")
        cells, p = [], 0
        for r in rows:
            rc = []
            for c in r:
                rc.append((c, start + p))
                p += len(c) + 1
            cells.append(rc)
        blocks.append({"type": "table", "start": start, "cells": cells, "align": []})
        offset += len(txt) + 1

    def emit_picture(shape):
        nonlocal img_counter
        try:
            blob = shape.image.blob
        except Exception:
            return
        img_id = f"pptx_img_{img_counter}"
        img_counter += 1
        images[img_id] = blob
        w = h = 0
        try:
            from PIL import Image as PILImage
            with PILImage.open(io.BytesIO(blob)) as im:
                w, h = im.size
        except Exception:
            pass
        blocks.append({"type": "image", "start": offset, "img_id": img_id,
                       "img_w": w, "img_h": h})

    def emit_text_frame(shape):
        ph = _ph_type(shape)
        in_body = ph in (PH.BODY, PH.OBJECT)
        counters = {}
        for p in shape.text_frame.paragraphs:
            pieces = _para_pieces(p, offset)
            if not "".join(x[0] for x in pieces).strip():
                continue
            kind = _bullet_kind(p, in_body)
            depth = p.level
            if kind is None:
                counters = {}
                emit_text("para", pieces)
                continue
            if kind == "num":
                counters[depth] = counters.get(depth, 0) + 1
                for d in [d for d in counters if d > depth]:
                    del counters[d]
                marker = f"{counters[depth]}."
            else:
                marker = "•"
            emit_text("list", pieces, marker=marker, depth=depth)

    total = max(1, len(prs.slides))
    for n, slide in enumerate(prs.slides, 1):
        shapes = list(_iter_shapes(slide.shapes))
        stitle = ""
        for s in shapes:
            if _is_title(s) and s.has_text_frame:
                stitle = " ".join(s.text_frame.text.replace("\x0b", " ").split())
                if stitle:
                    break
        stitle = stitle or f"Slide {n}"
        start, _ = emit_text("heading", [(stitle, 0, offset, offset + len(stitle))], level=1)
        chapters.append((start, stitle))

        for sh in shapes:
            if _is_title(sh):
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                emit_table(sh.table)
            elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(sh, "image"):
                emit_picture(sh)
            elif sh.has_text_frame:
                emit_text_frame(sh)
        if progress:
            progress(5 + int(90 * n / total), f"解析第 {n} 页…")

    text = "".join(text_parts)
    if not chapters:
        from .textio import scan_chapters
        chapters = scan_chapters(text)
    if progress:
        progress(95, "完成")
    return text, chapters, title, {"blocks": blocks, "images": images}
